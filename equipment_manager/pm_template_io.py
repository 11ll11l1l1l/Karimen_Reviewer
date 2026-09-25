from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from services import parse_spec


DEFINITION_HEADERS = [
    "PM ID","Name","Equipment ID","Schedule Type","Frequency","Unit","Anchor",
    "Early Window Days","Grace Days","Estimated Hours","Required People",
    "Required Skill","Required Parts","Primary Instruction Path",
]
STEP_HEADERS = [
    "Step","Activity / Instruction","Method","Input Type","Unit","Target",
    "Warning Low","Warning High","Control Low","Control High","Spec Low","Spec High",
    "Acceptance","Reaction Plan","Screenshot Required","Comment Required",
    "Reference Path","Reference Page","Reference Section",
]
REQUIREMENT_HEADERS = [
    "Requirement ID","Type","Key / Code","Description","Quantity","Mandatory",
]


def _style_sheet(ws):
    fill=PatternFill("solid",fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font=Font(bold=True)
        cell.fill=fill
    ws.freeze_panes="A2"
    for col in range(1,ws.max_column+1):
        letter=get_column_letter(col)
        width=max(12,min(46,max((len(str(ws.cell(r,col).value or "")) for r in range(1,ws.max_row+1)),default=12)+2))
        ws.column_dimensions[letter].width=width


def generate_pm_template(path: str, definition=None, specs=None, requirements=None) -> str:
    """Generate an EMS-authored workbook that users can edit and re-import."""
    specs=list(specs or [])
    requirements=list(requirements or [])
    wb=Workbook()
    readme=wb.active
    readme.title="README"
    readme.append(["EMS PM TEMPLATE","Instructions"])
    guidance=[
        ("Purpose","Edit this workbook using the content from the current PM Excel/SOP, then import it into EMS."),
        ("PM Definition","One PM definition per workbook. Keep the PM ID stable when revising an existing PM."),
        ("Steps","Add or remove rows as needed. Step numbers must be unique within the PM."),
        ("Instructions","Put the actual step instruction in Activity / Instruction. Reference Path may point to an Excel, PowerPoint, PDF, or other controlled file."),
        ("Acceptance","Use Input Type = Text, Numeric, or Pass / Fail. Numeric limits may be entered directly."),
        ("Completion evidence","Set Screenshot Required and/or Comment Required to Yes when that step must include that proof before PM completion."),
        ("Requirements","Use CERTIFICATION, LOTO, SAFETY, TOOL, PART, or DOCUMENT."),
        ("Unknown sheets","Extra sheets are allowed and ignored by the importer, so supporting material may remain in the workbook."),
        ("Images","Images may remain in supporting sheets. For controlled execution, reference the source document/path from the step."),
        ("Revision safety","Import previews the PM ID and row counts before applying. Existing PM steps are revised rather than silently overwritten."),
    ]
    for row in guidance:readme.append(row)
    readme.column_dimensions["A"].width=24;readme.column_dimensions["B"].width=100
    readme["A1"].font=Font(bold=True);readme["B1"].font=Font(bold=True)

    d=wb.create_sheet("PM Definition");d.append(DEFINITION_HEADERS)
    if definition:
        d.append([
            definition.pm_id,definition.name,definition.equipment_id,definition.schedule_type,
            definition.frequency_value,definition.frequency_unit,definition.anchor_mode,
            definition.early_window_days,definition.grace_days,definition.estimated_hours,
            definition.required_people,definition.required_skill,definition.required_parts,
            definition.sop_path,
        ])
    else:
        d.append(["","","","Interval",1,"months","Original Due",0,0,1,1,"","",""])
    _style_sheet(d)

    s=wb.create_sheet("Steps");s.append(STEP_HEADERS)
    for row in specs:
        s.append([
            row.step_no,row.activity,row.method,row.input_type,row.unit,row.target,
            row.warning_low,row.warning_high,row.control_low,row.control_high,row.spec_low,row.spec_high,
            row.acceptance_text,row.reaction_plan,
            "Yes" if getattr(row,"screenshot_required",False) else "No",
            "Yes" if getattr(row,"comment_required",False) else "No",
            row.sop_path,row.sop_page,row.sop_section,
        ])
    if not specs:
        s.append([1,"Describe the work step","","Text","","","","","","","","","","","No","No","","",""])
    _style_sheet(s)

    r=wb.create_sheet("Requirements");r.append(REQUIREMENT_HEADERS)
    for row in requirements:
        r.append([
            row.requirement_id,row.requirement_type,row.requirement_key,row.description,
            row.quantity,"Yes" if row.mandatory else "No",
        ])
    _style_sheet(r)

    Path(path).parent.mkdir(parents=True,exist_ok=True)
    wb.save(path)
    return str(path)


def _records(ws) -> list[dict[str,Any]]:
    rows=list(ws.iter_rows(values_only=True))
    if not rows:return []
    headers=[str(x or "").strip() for x in rows[0]]
    result=[]
    for values in rows[1:]:
        if not any(v not in (None,"") for v in values):continue
        result.append({headers[i]:values[i] if i<len(values) else None for i in range(len(headers))})
    return result


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _float(value):
    if value in (None,""):return None
    try:return float(value)
    except Exception:return None


def _int(value,default=0):
    if value in (None,""):return default
    try:return int(float(value))
    except Exception:return default


def _bool(value,default=True):
    if value is None:return default
    if isinstance(value,bool):return value
    return _text(value).lower() not in {"0","false","no","n","off"}


def load_pm_template(path: str) -> dict[str,Any]:
    wb=load_workbook(path,data_only=True,read_only=False)
    required={"PM Definition","Steps"}
    missing=sorted(required-set(wb.sheetnames))
    if missing:raise ValueError("PM template is missing sheet(s): "+", ".join(missing))
    warnings=[]

    defs=_records(wb["PM Definition"])
    if not defs:raise ValueError("PM Definition sheet has no definition row.")
    if len(defs)>1:warnings.append("Only the first PM Definition row is imported.")
    x=defs[0]
    definition={
        "pm_id":_text(x.get("PM ID")),
        "name":_text(x.get("Name")),
        "equipment_id":_text(x.get("Equipment ID")),
        "schedule_type":_text(x.get("Schedule Type")) or "Interval",
        "frequency_value":_float(x.get("Frequency")),
        "frequency_unit":_text(x.get("Unit")) or "days",
        "anchor_mode":_text(x.get("Anchor")) or "Original Due",
        "early_window_days":_int(x.get("Early Window Days"),0),
        "grace_days":_int(x.get("Grace Days"),0),
        "estimated_hours":_float(x.get("Estimated Hours")) or 0.0,
        "required_people":max(1,_int(x.get("Required People"),1)),
        "required_skill":_text(x.get("Required Skill")),
        "required_parts":_text(x.get("Required Parts")),
        "sop_path":_text(x.get("Primary Instruction Path")),
        "active":True,
    }
    if not definition["pm_id"]:raise ValueError("PM ID is required.")
    if not definition["name"]:definition["name"]=definition["pm_id"]

    steps=[]
    seen_steps=set()
    for index,row in enumerate(_records(wb["Steps"]),start=2):
        step_no=_int(row.get("Step"),0)
        activity=_text(row.get("Activity / Instruction"))
        if step_no<=0 or not activity:
            warnings.append(f"Steps row {index}: Step and Activity / Instruction are required; row ignored.")
            continue
        if step_no in seen_steps:
            warnings.append(f"Steps row {index}: duplicate step {step_no}; row ignored.")
            continue
        seen_steps.add(step_no)
        input_type=_text(row.get("Input Type"))
        acceptance=_text(row.get("Acceptance"))
        parsed=parse_spec(acceptance) if acceptance and not input_type else {}
        spec={
            "pm_id":definition["pm_id"],"step_no":step_no,"activity":activity,
            "method":_text(row.get("Method")),
            "input_type":input_type or parsed.get("input_type","Text"),
            "unit":_text(row.get("Unit")),
            "target":_float(row.get("Target")),
            "warning_low":_float(row.get("Warning Low")),"warning_high":_float(row.get("Warning High")),
            "control_low":_float(row.get("Control Low")),"control_high":_float(row.get("Control High")),
            "spec_low":_float(row.get("Spec Low")),"spec_high":_float(row.get("Spec High")),
            "acceptance_text":acceptance or _text(parsed.get("acceptance_text","")),
            "reaction_plan":_text(row.get("Reaction Plan")),
            "screenshot_required":_bool(row.get("Screenshot Required"),False),
            "comment_required":_bool(row.get("Comment Required"),False),
            "sop_path":_text(row.get("Reference Path")),
            "sop_page":_text(row.get("Reference Page")),
            "sop_section":_text(row.get("Reference Section")),
            "active":True,
        }
        for key in ["target","spec_low","spec_high"]:
            if spec[key] is None and key in parsed:spec[key]=parsed[key]
        if spec["spec_low"] is not None and spec["spec_high"] is not None and spec["spec_low"]>spec["spec_high"]:
            warnings.append(f"Steps row {index}: lower spec exceeds upper spec; row ignored.")
            continue
        steps.append(spec)

    requirements=[]
    if "Requirements" in wb.sheetnames:
        for index,row in enumerate(_records(wb["Requirements"]),start=2):
            rid=_text(row.get("Requirement ID"));rtype=_text(row.get("Type")).upper()
            desc=_text(row.get("Description"))
            if not rid and not rtype and not desc:continue
            if not rid:
                rid=f"{definition['pm_id']}-REQ-{index-1:03d}"
                warnings.append(f"Requirements row {index}: generated requirement ID {rid}.")
            if rtype not in {"CERTIFICATION","LOTO","SAFETY","TOOL","PART","DOCUMENT"}:
                warnings.append(f"Requirements row {index}: unsupported type {rtype or '(blank)'}; row ignored.")
                continue
            requirements.append({
                "requirement_id":rid,"pm_id":definition["pm_id"],"requirement_type":rtype,
                "requirement_key":_text(row.get("Key / Code")),"description":desc or rid,
                "quantity":_float(row.get("Quantity")) or 1.0,
                "mandatory":_bool(row.get("Mandatory"),True),"active":True,
            })

    extras=[x for x in wb.sheetnames if x not in {"README","PM Definition","Steps","Requirements"}]
    if extras:warnings.append("Supporting sheet(s) retained but not imported: "+", ".join(extras))
    return {"definition":definition,"steps":steps,"requirements":requirements,"warnings":warnings}
