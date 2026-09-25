from __future__ import annotations

import csv
import html
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def _status_counts(rows):
    return {
        "total":len(rows),
        "good":sum(1 for x in rows if x.get("health")=="good"),
        "critical":sum(1 for x in rows if x.get("health")=="critical"),
        "attention":sum(1 for x in rows if x.get("health")=="attention"),
        "planned":sum(1 for x in rows if x.get("health")=="planned"),
        "offline":sum(1 for x in rows if x.get("health")=="offline"),
        "alarms":sum(int(x.get("active_alarm_count") or 0) for x in rows),
        "tickets":sum(int(x.get("open_ticket_count") or 0) for x in rows),
    }


def report_rows(rows: list[dict[str,Any]]) -> list[dict[str,Any]]:
    out=[]
    for x in rows:
        if x.get("health")=="good" and not x.get("next_pm_within_24h") and not x.get("open_work_order_count"):
            continue
        out.append({
            "equipment_id":x.get("equipment_id",""),
            "state":x.get("equipment_status",""),
            "health":x.get("health",""),
            "issue":x.get("current_issue_title",""),
            "lots":", ".join(x.get("lots") or []),
            "alarms":", ".join(x.get("alarm_codes") or []),
            "owner":x.get("owner",""),
            "action":x.get("current_action",""),
            "active_pm":x.get("active_pm_id",""),
            "next_pm":x.get("next_pm_id",""),
            "next_pm_at":x.get("next_pm_at"),
        })
    rank={"critical":0,"attention":1,"planned":2,"offline":3,"good":4}
    out.sort(key=lambda x:(rank.get(x["health"],9),x["equipment_id"]))
    return out


def build_fab_status_report(
    rows: list[dict[str,Any]],
    *,
    title: str="FAB Status",
    generated_at: datetime | None=None,
) -> tuple[str,str]:
    generated_at=generated_at or datetime.now()
    counts=_status_counts(rows);details=report_rows(rows)
    summary=(
        f"{counts['total']} tools | {counts['good']} Running/Good | "
        f"{counts['critical']} Critical/Down | {counts['attention']} Attention | "
        f"{counts['planned']} PM/Engineering | {counts['offline']} Offline | "
        f"{counts['alarms']} Active Alarms | {counts['tickets']} Open Issues"
    )
    plain=[f"{title} — {generated_at:%Y-%m-%d %H:%M}",summary,""]
    for group,label in [("critical","CRITICAL / DOWN"),("attention","ATTENTION"),("planned","PM / ENGINEERING"),("offline","OFFLINE"),("good","UPCOMING / OTHER")]:
        grouped=[x for x in details if x["health"]==group]
        if not grouped:continue
        plain.append(label)
        for x in grouped:
            bits=[x["equipment_id"],x["state"]]
            if x["issue"]:bits.append(x["issue"])
            if x["lots"]:bits.append("Lot "+x["lots"])
            if x["alarms"]:bits.append("Alarm "+x["alarms"])
            if x["action"]:bits.append("Action: "+x["action"])
            if x["owner"]:bits.append("Owner: "+x["owner"])
            if x["active_pm"]:bits.append("PM: "+x["active_pm"])
            elif x["next_pm"] and x["next_pm_at"]:bits.append(f"Next PM: {x['next_pm']} {x['next_pm_at']:%m-%d %H:%M}")
            plain.append(" - "+" | ".join(bits))
        plain.append("")
    plain_text="\n".join(plain).rstrip()

    rows_html=[]
    for x in details:
        next_pm=""
        if x["active_pm"]:next_pm="Active "+x["active_pm"]
        elif x["next_pm"]:
            next_pm=x["next_pm"]+(f" {x['next_pm_at']:%m-%d %H:%M}" if x["next_pm_at"] else "")
        rows_html.append(
            "<tr>"
            +f"<td>{html.escape(str(x['equipment_id']))}</td>"
            +f"<td>{html.escape(str(x['state']))}</td>"
            +f"<td>{html.escape(str(x['issue']))}</td>"
            +f"<td>{html.escape(str(x['lots']))}</td>"
            +f"<td>{html.escape(str(x['alarms']))}</td>"
            +f"<td>{html.escape(str(x['action']))}</td>"
            +f"<td>{html.escape(str(x['owner']))}</td>"
            +f"<td>{html.escape(str(next_pm))}</td>"
            +"</tr>"
        )
    html_text=f"""<div style="font-family:Segoe UI,Arial,sans-serif">
<h2 style="margin-bottom:4px">{html.escape(title)}</h2>
<div style="color:#555">{generated_at:%Y-%m-%d %H:%M}</div>
<p><b>{html.escape(summary)}</b></p>
<table style="border-collapse:collapse;font-size:10pt" border="1" cellpadding="5">
<tr style="background:#e9edf1"><th>Tool</th><th>State</th><th>Issue</th><th>Lot(s)</th><th>Alarm(s)</th><th>Current Action</th><th>Owner</th><th>PM</th></tr>
{''.join(rows_html)}
</table></div>"""
    return plain_text,html_text


def export_fab_status_csv(rows: list[dict[str,Any]],path: str) -> str:
    details=report_rows(rows)
    fields=["equipment_id","state","health","issue","lots","alarms","owner","action","active_pm","next_pm","next_pm_at"]
    Path(path).parent.mkdir(parents=True,exist_ok=True)
    with open(path,"w",encoding="utf-8-sig",newline="") as handle:
        writer=csv.DictWriter(handle,fieldnames=fields);writer.writeheader()
        for row in details:
            data=dict(row);data["next_pm_at"]=data["next_pm_at"].isoformat(sep=" ") if data["next_pm_at"] else ""
            writer.writerow(data)
    return str(path)


def export_fab_status_xlsx(rows: list[dict[str,Any]],path: str) -> str:
    wb=Workbook();ws=wb.active;ws.title="FAB Status"
    counts=_status_counts(rows)
    ws.append(["FAB STATUS",datetime.now().strftime("%Y-%m-%d %H:%M")])
    ws.append(["Total",counts["total"],"Running / Good",counts["good"],"Critical / Down",counts["critical"],"Attention",counts["attention"],"PM / Engineering",counts["planned"],"Offline",counts["offline"]])
    headers=["Tool","State","Health","Issue","Lots","Alarms","Owner","Current Action","Active PM","Next PM","Next PM At"]
    ws.append([]);ws.append(headers)
    for cell in ws[4]:
        cell.font=Font(bold=True);cell.fill=PatternFill("solid",fgColor="D9EAF7")
    for x in report_rows(rows):
        ws.append([x["equipment_id"],x["state"],x["health"],x["issue"],x["lots"],x["alarms"],x["owner"],x["action"],x["active_pm"],x["next_pm"],x["next_pm_at"]])
    for col in ws.columns:
        letter=col[0].column_letter
        ws.column_dimensions[letter].width=min(48,max(12,max(len(str(c.value or "")) for c in col)+2))
    Path(path).parent.mkdir(parents=True,exist_ok=True);wb.save(path);return str(path)


def export_fab_status_pdf(rows: list[dict[str,Any]],path: str,title: str="FAB Status") -> str:
    counts=_status_counts(rows);styles=getSampleStyleSheet();story=[]
    story.append(Paragraph(title,styles["Heading1"]))
    story.append(Paragraph(datetime.now().strftime("%Y-%m-%d %H:%M"),styles["BodyText"]))
    story.append(Paragraph(
        f"{counts['total']} tools · {counts['good']} Running/Good · {counts['critical']} Critical/Down · "
        f"{counts['attention']} Attention · {counts['planned']} PM/Engineering · "
        f"{counts['alarms']} Active Alarms · {counts['tickets']} Open Issues",styles["BodyText"]
    ));story.append(Spacer(1,10))
    data=[["Tool","State","Issue","Lot(s)","Alarm(s)","Action","Owner","PM"]]
    for x in report_rows(rows):
        pm=x["active_pm"] or (x["next_pm"]+(f" {x['next_pm_at']:%m-%d %H:%M}" if x["next_pm_at"] else "") if x["next_pm"] else "")
        data.append([x["equipment_id"],x["state"],x["issue"],x["lots"],x["alarms"],x["action"],x["owner"],pm])
    table=Table(data,repeatRows=1,colWidths=[65,55,120,75,75,150,70,95])
    table.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#D9EAF7")),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
        ("FONTSIZE",(0,0),(-1,-1),6.8),("VALIGN",(0,0),(-1,-1),"TOP"),
        ("GRID",(0,0),(-1,-1),.35,colors.HexColor("#B8C3CB")),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F7F9FA")]),
    ]));story.append(table)
    Path(path).parent.mkdir(parents=True,exist_ok=True)
    doc=SimpleDocTemplate(path,pagesize=landscape(A4),leftMargin=20,rightMargin=20,topMargin=22,bottomMargin=22)
    doc.build(story);return str(path)
