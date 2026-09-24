from __future__ import annotations

from datetime import datetime
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.util import Inches, Pt


def _text(value: Any) -> str:
    if value is None:return ""
    if isinstance(value,datetime):return value.strftime("%Y-%m-%d %H:%M")
    return str(value)


def _ticket(db,ticket_no: str):
    row=next((x for x in db.list_tickets() if x.ticket_no==ticket_no),None)
    if not row:raise ValueError("Ticket not found")
    return row


def _set_title(slide,title: str):
    if slide.shapes.title:
        slide.shapes.title.text=title
        slide.shapes.title.text_frame.paragraphs[0].font.size=Pt(26)


def _add_bullets(slide,title: str,lines: list[str]):
    _set_title(slide,title)
    box=slide.placeholders[1] if len(slide.placeholders)>1 else slide.shapes.add_textbox(Inches(.8),Inches(1.5),Inches(11.5),Inches(5.5))
    tf=box.text_frame;tf.clear()
    for i,line in enumerate([x for x in lines if x]):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.level=0;p.font.size=Pt(18)


def _add_table(slide,title: str,headers: list[str],rows: list[list[Any]],max_rows: int=14):
    _set_title(slide,title)
    rows=rows[:max_rows]
    shape=slide.shapes.add_table(len(rows)+1,len(headers),Inches(.45),Inches(1.45),Inches(12.4),Inches(5.65))
    table=shape.table
    for c,h in enumerate(headers):
        cell=table.cell(0,c);cell.text=str(h)
        for p in cell.text_frame.paragraphs:p.font.bold=True;p.font.size=Pt(12)
    for r,row in enumerate(rows,1):
        for c,val in enumerate(row):
            table.cell(r,c).text=_text(val)
            for p in table.cell(r,c).text_frame.paragraphs:p.font.size=Pt(10)


def _add_evidence_slides(prs: Presentation,attachments,title_prefix: str):
    images=[x for x in attachments if (x.media_type or "").startswith("image/") and Path(x.stored_path).is_file()]
    for index,row in enumerate(images[:6],1):
        slide=prs.slides.add_slide(prs.slide_layouts[5]);_set_title(slide,f"{title_prefix} Evidence {index}")
        slide.shapes.add_picture(row.stored_path,Inches(.7),Inches(1.4),width=Inches(7.7),height=Inches(5.2))
        tb=slide.shapes.add_textbox(Inches(8.7),Inches(1.5),Inches(4),Inches(4.8))
        tf=tb.text_frame;tf.text=row.caption or row.original_name
        p=tf.add_paragraph();p.text=f"Category: {row.category}"
        p=tf.add_paragraph();p.text=f"Added by: {row.created_by}"
        if row.tags:
            p=tf.add_paragraph();p.text=f"Tags: {row.tags}"


def export_incident_pptx(db,ticket_no: str,path: str) -> str:
    t=_ticket(db,ticket_no);control=db.ticket_operational_control(ticket_no)
    whys=db.list_incident_whys(ticket_no);factors=db.list_incident_causal_factors(ticket_no)
    actions=db.list_incident_actions(ticket_no);lifecycle=db.list_ticket_state_events(ticket_no)
    attachments=db.list_attachments("TICKET",ticket_no)

    prs=Presentation()
    slide=prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text=f"{t.ticket_no} — {t.title}"
    slide.placeholders[1].text=f"{t.equipment_id} | {t.priority} | {t.status} | Owner: {t.owner or '—'}"

    slide=prs.slides.add_slide(prs.slide_layouts[1])
    _add_bullets(slide,"Problem / Operational Impact",[
        f"Problem: {t.description}",
        f"Production impact: {getattr(control,'production_impact','') if control else ''}",
        f"Affected lots/material: {getattr(control,'affected_lots','') if control else ''}",
        f"Safety/quality risk: {getattr(control,'safety_quality_risk','') if control else ''}",
        f"Containment: {getattr(control,'containment','') if control else ''}",
    ])

    slide=prs.slides.add_slide(prs.slide_layouts[5])
    _add_table(slide,"5-Why / Causal Analysis",["#","Question","Answer"],[[x.sequence,x.question,x.answer] for x in whys],10)

    slide=prs.slides.add_slide(prs.slide_layouts[5])
    _add_table(slide,"Causal Factors",["Category","Type","Description","Evidence","Status"],[
        [x.category,x.factor_type,x.description,x.evidence,x.status] for x in factors
    ],10)

    slide=prs.slides.add_slide(prs.slide_layouts[5])
    _add_table(slide,"Corrective / Preventive Actions",["Type","Action","Owner","Due","Status","Effectiveness"],[
        [x.action_type,x.description,x.owner,x.due_at,x.status,x.effectiveness_criteria] for x in actions
    ],12)

    slide=prs.slides.add_slide(prs.slide_layouts[1])
    _add_bullets(slide,"Resolution / Verification",[
        f"Root cause summary: {t.root_cause or 'Open'}",
        f"Corrective action summary: {t.corrective_action or 'Open'}",
        f"Verification: {t.verification or 'Pending'}",
        f"Current status: {t.status}",
    ])

    slide=prs.slides.add_slide(prs.slide_layouts[5])
    _add_table(slide,"Lifecycle Timeline",["From","To","Reason","Owner","By","Time"],[
        [x.from_state,x.to_state,x.reason_code,x.owner,x.changed_by,x.changed_at] for x in lifecycle
    ],16)

    _add_evidence_slides(prs,attachments,"Incident")
    Path(path).parent.mkdir(parents=True,exist_ok=True);prs.save(path);return str(path)


def export_equipment_pptx(db,equipment_id: str,path: str) -> str:
    eq=db.get_equipment(equipment_id)
    if not eq:raise ValueError("Equipment not found")
    rel=db.reliability_summary(equipment_id)
    tickets=[x for x in db.list_tickets() if x.equipment_id==equipment_id and x.status not in {"Closed","Cancelled"}]
    pm=[x for x in db.list_pm_tasks() if x.equipment_id==equipment_id and x.status not in {"Completed","Cancelled"}]
    wo=db.list_work_orders(equipment_id,True)
    alarms=db.list_alarms(equipment_id,True,100)
    timeline=db.equipment_activity_timeline(equipment_id,80)
    attachments=db.list_attachments("EQUIPMENT",equipment_id)

    prs=Presentation()
    slide=prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text=f"{eq.equipment_id} — {eq.name}"
    slide.placeholders[1].text=f"{eq.status} | {eq.disposition} | {eq.area} | Owner: {eq.owner or '—'}"

    slide=prs.slides.add_slide(prs.slide_layouts[1])
    _add_bullets(slide,"Current Equipment Summary",[
        f"Equipment type: {eq.equipment_type}",
        f"Model / serial: {eq.model} / {eq.serial_number}",
        f"Location: {eq.site} / {eq.building} / {eq.floor} / {eq.area} / {eq.line_cell}",
        f"Availability (30d): {rel['availability_pct']:.1f}%",
        f"MTBF (30d): {rel['mtbf_hours']:.1f} h",
        f"MTTR (30d): {rel['mttr_hours']:.1f} h",
        f"Unplanned downtime (30d): {rel['unplanned_downtime_hours']:.1f} h",
    ])

    slide=prs.slides.add_slide(prs.slide_layouts[5])
    _add_table(slide,"Open Incidents",["Ticket","Priority","Status","Title","Owner"],[[x.ticket_no,x.priority,x.status,x.title,x.owner] for x in tickets],12)

    slide=prs.slides.add_slide(prs.slide_layouts[5])
    _add_table(slide,"Maintenance / Work Orders",["Type","Key","Status","Owner / Assigned","Summary"],[
        *[["PM",x.id,x.status,x.assigned_to,f"{x.pm_id} — {x.pm_name}"] for x in pm],
        *[["WO",x.work_order_no,x.status,x.owner,x.title] for x in wo],
    ],14)

    slide=prs.slides.add_slide(prs.slide_layouts[5])
    _add_table(slide,"Active Alarms",["Code","Severity","Message","State","Occurred"],[[x.alarm_code,x.severity,x.message,x.state,x.occurred_at] for x in alarms],14)

    slide=prs.slides.add_slide(prs.slide_layouts[5])
    _add_table(slide,"Recent Activity",["Time","Type","Key","Activity","Status"],[
        [x["occurred_at"],x["kind"],x["key"],x["summary"],x["status"]] for x in timeline
    ],18)

    _add_evidence_slides(prs,attachments,"Equipment")
    Path(path).parent.mkdir(parents=True,exist_ok=True);prs.save(path);return str(path)


def _sheet(ws,headers: list[str],rows: list[list[Any]]):
    for c,h in enumerate(headers,1):
        cell=ws.cell(1,c,h);cell.font=Font(bold=True);cell.fill=PatternFill("solid",fgColor="D9E2F3")
    for r,row in enumerate(rows,2):
        for c,val in enumerate(row,1):ws.cell(r,c,_text(val))
    ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions
    for c,h in enumerate(headers,1):
        width=max([len(h)]+[len(_text(ws.cell(r,c).value)) for r in range(2,min(ws.max_row,300)+1)])
        ws.column_dimensions[get_column_letter(c)].width=min(max(width+2,10),60)


def export_incident_xlsx(db,ticket_no: str,path: str) -> str:
    t=_ticket(db,ticket_no);control=db.ticket_operational_control(ticket_no)
    wb=Workbook();summary=wb.active;summary.title="Summary"
    rows=[
        ["Ticket",t.ticket_no],["Equipment",t.equipment_id],["Title",t.title],["Priority",t.priority],["Status",t.status],
        ["Owner",t.owner],["Problem",t.description],["Containment",getattr(control,"containment","") if control else ""],
        ["Production impact",getattr(control,"production_impact","") if control else ""],["Affected lots",getattr(control,"affected_lots","") if control else ""],
        ["Root cause",t.root_cause],["Corrective action",t.corrective_action],["Verification",t.verification],
    ]
    _sheet(summary,["Field","Value"],rows)
    ws=wb.create_sheet("5-Why");_sheet(ws,["#","Question","Answer","Updated By","Updated"],[[x.sequence,x.question,x.answer,x.updated_by,x.updated_at] for x in db.list_incident_whys(ticket_no)])
    ws=wb.create_sheet("Causal Factors");_sheet(ws,["Category","Type","Description","Evidence","Status"],[[x.category,x.factor_type,x.description,x.evidence,x.status] for x in db.list_incident_causal_factors(ticket_no)])
    ws=wb.create_sheet("CAPA");_sheet(ws,["Type","Action","Owner","Due","Status","Effectiveness","Completed","Verified"],[[x.action_type,x.description,x.owner,x.due_at,x.status,x.effectiveness_criteria,x.completed_at,x.verified_at] for x in db.list_incident_actions(ticket_no)])
    ws=wb.create_sheet("Lifecycle");_sheet(ws,["From","To","Reason","Note","Owner","By","Time"],[[x.from_state,x.to_state,x.reason_code,x.note,x.owner,x.changed_by,x.changed_at] for x in db.list_ticket_state_events(ticket_no)])
    ws=wb.create_sheet("Attachments");_sheet(ws,["Name","Category","Caption","Tags","Path","Added By","Added"],[[x.original_name,x.category,x.caption,x.tags,x.stored_path,x.created_by,x.created_at] for x in db.list_attachments("TICKET",ticket_no)])
    Path(path).parent.mkdir(parents=True,exist_ok=True);wb.save(path);return str(path)


def export_equipment_xlsx(db,equipment_id: str,path: str) -> str:
    eq=db.get_equipment(equipment_id)
    if not eq:raise ValueError("Equipment not found")
    wb=Workbook();summary=wb.active;summary.title="Summary"
    rel=db.reliability_summary(equipment_id)
    _sheet(summary,["Field","Value"],[
        ["Equipment",eq.equipment_id],["Name",eq.name],["Status",eq.status],["Disposition",eq.disposition],
        ["Area",eq.area],["Line/Cell",eq.line_cell],["Owner",eq.owner],["Criticality",eq.criticality],
        ["Availability 30d",f"{rel['availability_pct']:.2f}%"],["MTBF 30d",rel["mtbf_hours"]],["MTTR 30d",rel["mttr_hours"]],
    ])
    timeline=db.equipment_activity_timeline(equipment_id,5000)
    ws=wb.create_sheet("Activity");_sheet(ws,["Time","Type","Key","Summary","User","Status","Source"],[[x["occurred_at"],x["kind"],x["key"],x["summary"],x["user"],x["status"],x["source"]] for x in timeline])
    tickets=[x for x in db.list_tickets() if x.equipment_id==equipment_id]
    ws=wb.create_sheet("Incidents");_sheet(ws,["Ticket","Title","Priority","Status","Owner","Created","Updated"],[[x.ticket_no,x.title,x.priority,x.status,x.owner,x.created_at,x.updated_at] for x in tickets])
    pm=[x for x in db.list_pm_tasks() if x.equipment_id==equipment_id]
    ws=wb.create_sheet("PM");_sheet(ws,["Task","PM","Name","Due","Scheduled","Status","Assigned","Hours"],[[x.id,x.pm_id,x.pm_name,x.original_due_date,x.scheduled_date,x.status,x.assigned_to,x.estimated_hours] for x in pm])
    wo=db.list_work_orders(equipment_id)
    ws=wb.create_sheet("Work Orders");_sheet(ws,["WO","Source","Title","Priority","Status","Owner","Created","Completed"],[[x.work_order_no,f"{x.source_type}:{x.source_key}",x.title,x.priority,x.status,x.owner,x.created_at,x.completed_at] for x in wo])
    alarms=db.list_alarms(equipment_id,False,5000)
    ws=wb.create_sheet("Alarms");_sheet(ws,["Code","Severity","Message","State","Occurred","Cleared","Ticket"],[[x.alarm_code,x.severity,x.message,x.state,x.occurred_at,x.cleared_at,x.related_ticket] for x in alarms])
    ws=wb.create_sheet("Attachments");_sheet(ws,["Name","Category","Caption","Tags","Path","Added By","Added"],[[x.original_name,x.category,x.caption,x.tags,x.stored_path,x.created_by,x.created_at] for x in db.list_attachments("EQUIPMENT",equipment_id)])
    Path(path).parent.mkdir(parents=True,exist_ok=True);wb.save(path);return str(path)


def export_pm_execution_pptx(db,task_id: int,path: str) -> str:
    task=db.get_pm_task(int(task_id))
    if not task:raise ValueError("PM task not found")
    execution=db.get_pm_execution_for_task(task.id)
    if not execution:raise ValueError("PM execution has not started.")
    specs=db.list_pm_execution_specs(execution.id);results={x.step_no:x for x in db.list_pm_results(execution.id)}
    requirements=db.list_pm_execution_requirements(execution.id);acks={x.requirement_id:x for x in db.list_pm_requirement_acks(execution.id)}
    reservations=[x for x in db.list_reservations() if x.pm_task_id==task.id]
    logs=[x for x in db.list_work_logs(task.equipment_id,False,1000) if x.entity_type=="PM_EXECUTION" and x.entity_key==str(execution.id)]
    attachments=db.list_attachments("PM_EXECUTION",str(execution.id))
    prs=Presentation()
    slide=prs.slides.add_slide(prs.slide_layouts[0]);slide.shapes.title.text=f"{task.pm_id} — {task.pm_name}";slide.placeholders[1].text=f"{task.equipment_id} | {task.status} | Assigned: {task.assigned_to or '—'}"
    slide=prs.slides.add_slide(prs.slide_layouts[1]);_add_bullets(slide,"PM Execution Summary",[
        f"Task ID: {task.id}",f"Equipment: {task.equipment_id}",f"Scheduled: {_text(task.scheduled_date or task.original_due_date)}",
        f"Execution status: {execution.status}",f"Started: {_text(execution.started_at)} by {execution.started_by or '—'}",
        f"Completed: {_text(execution.completed_at)} by {execution.completed_by or '—'}",
    ])
    rows=[]
    for spec in specs:
        result=results.get(spec.step_no)
        value=(result.value_text if result and result.value_text else (result.value_numeric if result else ""))
        rows.append([spec.step_no,spec.activity,spec.method,spec.unit,result.result if result else "OPEN",value,result.comment if result else ""])
    slide=prs.slides.add_slide(prs.slide_layouts[5]);_add_table(slide,"Checklist / Measurements",["Step","Activity","Method","Unit","Result","Value","Comment"],rows,16)
    req_rows=[]
    for req in requirements:
        ack=acks.get(req.requirement_id)
        req_rows.append([req.requirement_type,req.requirement_key,req.description,req.quantity,req.mandatory,ack.acknowledged_by if ack else ""])
    slide=prs.slides.add_slide(prs.slide_layouts[5]);_add_table(slide,"Requirements / Readiness",["Type","Key","Description","Qty","Mandatory","Acknowledged"],req_rows,14)
    slide=prs.slides.add_slide(prs.slide_layouts[5]);_add_table(slide,"Parts / Reservations",["Part","Location","Qty","Status","Reserved By","Time"],[[x.part_number,x.location_code,x.quantity,x.status,x.reserved_by,x.reserved_at] for x in reservations],14)
    slide=prs.slides.add_slide(prs.slide_layouts[5]);_add_table(slide,"Labor / Work Log",["User","Type","Start","End","Minutes","Status","Note"],[[x.username,x.work_type,x.started_at,x.ended_at,x.duration_minutes,x.status,x.note] for x in logs],14)
    _add_evidence_slides(prs,attachments,"PM")
    Path(path).parent.mkdir(parents=True,exist_ok=True);prs.save(path);return str(path)


def export_pm_execution_xlsx(db,task_id: int,path: str) -> str:
    task=db.get_pm_task(int(task_id))
    if not task:raise ValueError("PM task not found")
    execution=db.get_pm_execution_for_task(task.id)
    if not execution:raise ValueError("PM execution has not started.")
    specs=db.list_pm_execution_specs(execution.id);results={x.step_no:x for x in db.list_pm_results(execution.id)}
    requirements=db.list_pm_execution_requirements(execution.id);acks={x.requirement_id:x for x in db.list_pm_requirement_acks(execution.id)}
    reservations=[x for x in db.list_reservations() if x.pm_task_id==task.id]
    logs=[x for x in db.list_work_logs(task.equipment_id,False,1000) if x.entity_type=="PM_EXECUTION" and x.entity_key==str(execution.id)]
    attachments=db.list_attachments("PM_EXECUTION",str(execution.id))
    wb=Workbook();summary=wb.active;summary.title="Summary"
    _sheet(summary,["Field","Value"],[
        ["Task ID",task.id],["Equipment",task.equipment_id],["PM ID",task.pm_id],["PM Name",task.pm_name],["Task Status",task.status],
        ["Scheduled",task.scheduled_date or task.original_due_date],["Execution",execution.status],["Started By",execution.started_by],["Started",execution.started_at],
        ["Completed By",execution.completed_by],["Completed",execution.completed_at],
    ])
    ws=wb.create_sheet("Checklist");_sheet(ws,["Step","Activity","Method","Unit","Result","Value Text","Value Numeric","Comment","Entered By","Entered"],[
        [s.step_no,s.activity,s.method,s.unit,(results.get(s.step_no).result if results.get(s.step_no) else "OPEN"),
         (results.get(s.step_no).value_text if results.get(s.step_no) else ""),(results.get(s.step_no).value_numeric if results.get(s.step_no) else ""),
         (results.get(s.step_no).comment if results.get(s.step_no) else ""),(results.get(s.step_no).entered_by if results.get(s.step_no) else ""),
         (results.get(s.step_no).entered_at if results.get(s.step_no) else "")] for s in specs
    ])
    ws=wb.create_sheet("Requirements");_sheet(ws,["ID","Type","Key","Description","Qty","Mandatory","Acknowledged By","Acknowledged"],[
        [r.requirement_id,r.requirement_type,r.requirement_key,r.description,r.quantity,r.mandatory,(acks.get(r.requirement_id).acknowledged_by if acks.get(r.requirement_id) else ""),(acks.get(r.requirement_id).acknowledged_at if acks.get(r.requirement_id) else "")] for r in requirements
    ])
    ws=wb.create_sheet("Parts");_sheet(ws,["Part","Location","Qty","Status","Reserved By","Reserved"],[[x.part_number,x.location_code,x.quantity,x.status,x.reserved_by,x.reserved_at] for x in reservations])
    ws=wb.create_sheet("Labor");_sheet(ws,["User","Type","Start","End","Minutes","Status","Note"],[[x.username,x.work_type,x.started_at,x.ended_at,x.duration_minutes,x.status,x.note] for x in logs])
    ws=wb.create_sheet("Attachments");_sheet(ws,["Name","Category","Caption","Tags","Path","Added By","Added"],[[x.original_name,x.category,x.caption,x.tags,x.stored_path,x.created_by,x.created_at] for x in attachments])
    Path(path).parent.mkdir(parents=True,exist_ok=True);wb.save(path);return str(path)


def export_work_order_pptx(db,work_order_no: str,path: str) -> str:
    wo=db.get_work_order(work_order_no)
    if not wo:raise ValueError("Work order not found")
    events=db.list_work_order_events(work_order_no);links=db.list_work_order_links(work_order_no)
    logs=[x for x in db.list_work_logs(wo.equipment_id,False,1000) if x.entity_type=="WORK_ORDER" and x.entity_key==work_order_no]
    attachments=db.list_attachments("WORK_ORDER",work_order_no)
    close=db.work_order_closeout_status(work_order_no)
    prs=Presentation()
    slide=prs.slides.add_slide(prs.slide_layouts[0]);slide.shapes.title.text=f"{wo.work_order_no} — {wo.title}";slide.placeholders[1].text=f"{wo.equipment_id} | {wo.priority} | {wo.status} | Owner: {wo.owner or '—'}"
    slide=prs.slides.add_slide(prs.slide_layouts[1]);_add_bullets(slide,"Work Scope / Closeout",[
        f"Source: {wo.source_type}:{wo.source_key or '—'}",f"Description: {wo.description}",f"Team: {wo.team or '—'}",
        f"Qualification required: {'Yes' if wo.qualification_required else 'No'}",f"Release required: {'Yes' if wo.release_required else 'No'}",
        f"Labor entries: {close.get('labor_entries',0)}",f"Evidence attachments: {close.get('attachment_count',0)}",
        f"Blockers: {'; '.join(close.get('blockers',[])) or 'None'}",
    ])
    slide=prs.slides.add_slide(prs.slide_layouts[5]);_add_table(slide,"Linked Records",["Type","Key","Relation","Created By","Created"],[[x.entity_type,x.entity_key,x.relation,x.created_by,x.created_at] for x in links],14)
    slide=prs.slides.add_slide(prs.slide_layouts[5]);_add_table(slide,"Lifecycle",["From","To","Reason","Owner","Changed By","Time"],[[x.from_state,x.to_state,x.reason,x.owner,x.changed_by,x.occurred_at] for x in events],16)
    slide=prs.slides.add_slide(prs.slide_layouts[5]);_add_table(slide,"Labor",["User","Type","Start","End","Minutes","Status","Note"],[[x.username,x.work_type,x.started_at,x.ended_at,x.duration_minutes,x.status,x.note] for x in logs],14)
    _add_evidence_slides(prs,attachments,"Work Order")
    Path(path).parent.mkdir(parents=True,exist_ok=True);prs.save(path);return str(path)


def export_work_order_xlsx(db,work_order_no: str,path: str) -> str:
    wo=db.get_work_order(work_order_no)
    if not wo:raise ValueError("Work order not found")
    events=db.list_work_order_events(work_order_no);links=db.list_work_order_links(work_order_no)
    logs=[x for x in db.list_work_logs(wo.equipment_id,False,1000) if x.entity_type=="WORK_ORDER" and x.entity_key==work_order_no]
    attachments=db.list_attachments("WORK_ORDER",work_order_no);close=db.work_order_closeout_status(work_order_no)
    wb=Workbook();summary=wb.active;summary.title="Summary"
    _sheet(summary,["Field","Value"],[
        ["Work Order",wo.work_order_no],["Equipment",wo.equipment_id],["Title",wo.title],["Status",wo.status],["Priority",wo.priority],["Owner",wo.owner],["Team",wo.team],
        ["Source",f"{wo.source_type}:{wo.source_key}"],["Description",wo.description],["Qualification Required",wo.qualification_required],["Release Required",wo.release_required],
        ["Closeout Blockers","; ".join(close.get("blockers",[]))],
    ])
    ws=wb.create_sheet("Lifecycle");_sheet(ws,["From","To","Reason","Owner","Changed By","Time"],[[x.from_state,x.to_state,x.reason,x.owner,x.changed_by,x.occurred_at] for x in events])
    ws=wb.create_sheet("Links");_sheet(ws,["Type","Key","Relation","Created By","Created"],[[x.entity_type,x.entity_key,x.relation,x.created_by,x.created_at] for x in links])
    ws=wb.create_sheet("Labor");_sheet(ws,["User","Type","Start","End","Minutes","Status","Note"],[[x.username,x.work_type,x.started_at,x.ended_at,x.duration_minutes,x.status,x.note] for x in logs])
    ws=wb.create_sheet("Attachments");_sheet(ws,["Name","Category","Caption","Tags","Path","Added By","Added"],[[x.original_name,x.category,x.caption,x.tags,x.stored_path,x.created_by,x.created_at] for x in attachments])
    Path(path).parent.mkdir(parents=True,exist_ok=True);wb.save(path);return str(path)


def _qualification_run(db,run_no: str):
    row=next((x for x in db.list_qualification_runs() if x.run_no==run_no),None)
    if not row:raise ValueError("Qualification run not found")
    return row


def _release_request(db,release_id: int):
    row=next((x for x in db.list_release_requests() if x.id==int(release_id)),None)
    if not row:raise ValueError("Release request not found")
    return row


def export_qualification_pptx(db,run_no: str,path: str) -> str:
    run=_qualification_run(db,run_no)
    checks=json.loads(run.frozen_checks_json or "[]");results=json.loads(run.results_json or "{}")
    attachments=db.list_attachments("QUALIFICATION",run.run_no)
    prs=Presentation()
    slide=prs.slides.add_slide(prs.slide_layouts[0]);slide.shapes.title.text=f"Qualification — {run.run_no}";slide.placeholders[1].text=f"{run.equipment_id} | {run.protocol_id} R{run.protocol_revision} | {run.status}"
    slide=prs.slides.add_slide(prs.slide_layouts[1]);_add_bullets(slide,"Qualification Summary",[
        f"Protocol: {run.protocol_name}",f"Started: {_text(run.started_at)} by {run.started_by}",
        f"Submitted: {_text(run.submitted_at)} by {run.submitted_by or '—'}",f"Verified: {_text(run.verified_at)} by {run.verified_by or '—'}",
        f"Approved: {_text(run.approved_at)} by {run.approved_by or '—'}",f"Expires: {_text(run.expires_at) or 'No expiry'}",
        f"Conclusion: {run.conclusion or '—'}",
    ])
    rows=[]
    for check in checks:
        cid=str(check.get("check_id") or check.get("id") or check.get("name") or "")
        result=results.get(cid,{}) if isinstance(results,dict) else {}
        rows.append([cid,check.get("name") or check.get("description") or "",check.get("acceptance") or "",result.get("result",""),result.get("value",""),result.get("comment","")])
    slide=prs.slides.add_slide(prs.slide_layouts[5]);_add_table(slide,"Qualification Checks",["Check","Description","Acceptance","Result","Value","Comment"],rows,16)
    _add_evidence_slides(prs,attachments,"Qualification")
    Path(path).parent.mkdir(parents=True,exist_ok=True);prs.save(path);return str(path)


def export_qualification_xlsx(db,run_no: str,path: str) -> str:
    run=_qualification_run(db,run_no)
    checks=json.loads(run.frozen_checks_json or "[]");results=json.loads(run.results_json or "{}")
    wb=Workbook();summary=wb.active;summary.title="Summary"
    _sheet(summary,["Field","Value"],[
        ["Run",run.run_no],["Equipment",run.equipment_id],["Protocol",run.protocol_id],["Protocol Revision",run.protocol_revision],["Protocol Name",run.protocol_name],
        ["Status",run.status],["Started By",run.started_by],["Started",run.started_at],["Submitted By",run.submitted_by],["Submitted",run.submitted_at],
        ["Verified By",run.verified_by],["Verified",run.verified_at],["Approved By",run.approved_by],["Approved",run.approved_at],["Expires",run.expires_at],["Conclusion",run.conclusion],
    ])
    rows=[]
    for check in checks:
        cid=str(check.get("check_id") or check.get("id") or check.get("name") or "")
        result=results.get(cid,{}) if isinstance(results,dict) else {}
        rows.append([cid,check.get("name") or check.get("description") or "",check.get("acceptance") or "",result.get("result",""),result.get("value",""),result.get("comment",""),result.get("entered_by",""),result.get("entered_at","")])
    ws=wb.create_sheet("Checks");_sheet(ws,["Check","Description","Acceptance","Result","Value","Comment","Entered By","Entered"],rows)
    ws=wb.create_sheet("Attachments");_sheet(ws,["Name","Category","Caption","Tags","Path","Added By","Added"],[[x.original_name,x.category,x.caption,x.tags,x.stored_path,x.created_by,x.created_at] for x in db.list_attachments("QUALIFICATION",run.run_no)])
    Path(path).parent.mkdir(parents=True,exist_ok=True);wb.save(path);return str(path)


def export_release_pptx(db,release_id: int,path: str) -> str:
    rel=_release_request(db,release_id);checks=json.loads(rel.checks_json or "{}")
    attachments=db.list_attachments("RELEASE",str(rel.id))
    prs=Presentation()
    slide=prs.slides.add_slide(prs.slide_layouts[0]);slide.shapes.title.text=f"Equipment Release — {rel.equipment_id}";slide.placeholders[1].text=f"Release #{rel.id} | {rel.status} | Ticket: {rel.related_ticket or '—'}"
    slide=prs.slides.add_slide(prs.slide_layouts[1]);_add_bullets(slide,"Release Summary",[
        f"Requested: {_text(rel.requested_at)} by {rel.requested_by}",f"Verified: {_text(rel.verified_at)} by {rel.verified_by or '—'}",
        f"Approved: {_text(rel.approved_at)} by {rel.approved_by or '—'}",f"Notes: {rel.notes or '—'}",
    ])
    slide=prs.slides.add_slide(prs.slide_layouts[5]);_add_table(slide,"Release Checklist",["Check","Pass"],[[key,"PASS" if value else "FAIL"] for key,value in checks.items()],20)
    _add_evidence_slides(prs,attachments,"Release")
    Path(path).parent.mkdir(parents=True,exist_ok=True);prs.save(path);return str(path)


def export_release_xlsx(db,release_id: int,path: str) -> str:
    rel=_release_request(db,release_id);checks=json.loads(rel.checks_json or "{}")
    wb=Workbook();summary=wb.active;summary.title="Summary"
    _sheet(summary,["Field","Value"],[
        ["Release ID",rel.id],["Equipment",rel.equipment_id],["Related Ticket",rel.related_ticket],["Status",rel.status],["Notes",rel.notes],
        ["Requested By",rel.requested_by],["Requested",rel.requested_at],["Verified By",rel.verified_by],["Verified",rel.verified_at],["Approved By",rel.approved_by],["Approved",rel.approved_at],
    ])
    ws=wb.create_sheet("Checklist");_sheet(ws,["Check","Pass"],[[key,bool(value)] for key,value in checks.items()])
    ws=wb.create_sheet("Attachments");_sheet(ws,["Name","Category","Caption","Tags","Path","Added By","Added"],[[x.original_name,x.category,x.caption,x.tags,x.stored_path,x.created_by,x.created_at] for x in db.list_attachments("RELEASE",str(rel.id))])
    Path(path).parent.mkdir(parents=True,exist_ok=True);wb.save(path);return str(path)


def _review_metrics(db,days: int):
    data=db.engineering_analytics(days)
    tools=data["tool_matrix"];pm=data["pm"]
    fleet=(sum(float(x["availability_pct"]) for x in tools)/len(tools)) if tools else 100.0
    total_unplanned=sum(float(x["unplanned_downtime_hours"]) for x in tools)
    failures=sum(int(x["failure_count"]) for x in tools)
    critical=[x for x in db.list_tickets() if x.status not in {"Closed","Cancelled"} and x.priority in {"P1","P2"}]
    overdue=[x for x in db.list_pm_tasks() if x.status=="Overdue" or (x.status not in {"Completed","Cancelled"} and x.original_due_date and x.original_due_date<data["end"])]
    attention=db.operations_attention_queue(200)
    return data,{
        "fleet_availability":fleet,"unplanned_hours":total_unplanned,"failures":failures,
        "critical_incidents":critical,"overdue_pm":overdue,"attention":attention,"pm":pm,
    }


def _weekly_presentation(template_path: str=""):
    template=Path(template_path) if template_path else None
    if template and template.is_file():
        prs=Presentation(str(template))
        if len(prs.slide_layouts)<6:
            raise ValueError("PowerPoint template must provide standard title, content and blank layouts.")
        return prs
    return Presentation()


def export_weekly_review_pptx(db,path: str,days: int=7,template_path: str="") -> str:
    data,metrics=_review_metrics(db,days)
    prs=_weekly_presentation(template_path)
    slide=prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text="Equipment Engineering Weekly Review"
    slide.placeholders[1].text=f"{data['start']:%Y-%m-%d} to {data['end']:%Y-%m-%d} | Generated {datetime.now():%Y-%m-%d %H:%M}"

    slide=prs.slides.add_slide(prs.slide_layouts[1])
    pm=metrics["pm"]
    _add_bullets(slide,"Fleet / Maintenance Scorecard",[
        f"Fleet availability: {metrics['fleet_availability']:.1f}%",
        f"Unplanned downtime: {metrics['unplanned_hours']:.1f} h",
        f"Failure entries: {metrics['failures']}",
        f"Open P1/P2 incidents: {len(metrics['critical_incidents'])}",
        f"PM compliance: {pm['compliance_pct']:.1f}% ({pm['completed']}/{pm['due']} due tasks completed)",
        f"PM overdue: {pm['overdue']} · Deferred: {pm['deferred']}",
    ])

    slide=prs.slides.add_slide(prs.slide_layouts[5])
    _add_table(slide,"Chronic / High-Downtime Tools",
        ["Equipment","Avail %","Failures","Unplanned h","MTTR h","MTBF h","Open Inc","Alarms","State"],
        [[x["equipment_id"],f"{x['availability_pct']:.1f}",x["failure_count"],f"{x['unplanned_downtime_hours']:.1f}",
          f"{x['mttr_hours']:.1f}",f"{x['mtbf_hours']:.1f}",x["open_incidents"],x["active_alarms"],x["current_state"]]
         for x in data["tool_matrix"][:12]],12)

    slide=prs.slides.add_slide(prs.slide_layouts[5])
    _add_table(slide,"Alarm Pareto",["Alarm Code","Message","Count"],
               [[x["alarm_code"],x["message"],x["count"]] for x in data["alarm_pareto"][:15]],15)

    critical=sorted(metrics["critical_incidents"],key=lambda x:(0 if x.priority=="P1" else 1,x.created_at))
    slide=prs.slides.add_slide(prs.slide_layouts[5])
    _add_table(slide,"Open Critical Incidents",["Ticket","Equipment","Priority","Status","Owner","Title","Updated"],
               [[x.ticket_no,x.equipment_id,x.priority,x.status,x.owner,x.title,x.updated_at] for x in critical],14)

    overdue=sorted(metrics["overdue_pm"],key=lambda x:(x.original_due_date or datetime.max,x.equipment_id,x.pm_id))
    slide=prs.slides.add_slide(prs.slide_layouts[5])
    _add_table(slide,"PM Exceptions / Overdue",["Task","Equipment","PM","Due","Planned","Status","Assigned","Priority"],
               [[x.id,x.equipment_id,x.pm_id,x.original_due_date,x.scheduled_date,x.status,x.assigned_to,x.priority] for x in overdue],14)

    slide=prs.slides.add_slide(prs.slide_layouts[5])
    _add_table(slide,"Current Operational Priorities",["Severity","Type","Equipment","Key","Condition / Action","Owner","Age h"],
               [[x.get("severity"),x.get("kind"),x.get("equipment_id"),x.get("key"),x.get("summary"),x.get("owner"),f"{float(x.get('age_hours') or 0):.1f}"] for x in metrics["attention"][:16]],16)

    Path(path).parent.mkdir(parents=True,exist_ok=True);prs.save(path);return str(path)


def export_weekly_review_xlsx(db,path: str,days: int=7) -> str:
    data,metrics=_review_metrics(db,days);pm=metrics["pm"]
    wb=Workbook();summary=wb.active;summary.title="Scorecard"
    _sheet(summary,["Metric","Value"],[
        ["Period start",data["start"]],["Period end",data["end"]],
        ["Fleet availability %",f"{metrics['fleet_availability']:.2f}"],
        ["Unplanned downtime h",f"{metrics['unplanned_hours']:.2f}"],["Failures",metrics["failures"]],
        ["Open P1/P2 incidents",len(metrics["critical_incidents"])],["PM compliance %",f"{pm['compliance_pct']:.2f}"],
        ["PM due",pm["due"]],["PM completed",pm["completed"]],["PM overdue",pm["overdue"]],["PM deferred",pm["deferred"]],
    ])
    ws=wb.create_sheet("Chronic Tools");_sheet(ws,
        ["Equipment","Availability %","Failures","Unplanned h","Planned h","MTTR h","MTBF h","Incidents","Open","P1/P2 Open","Active Alarms","State"],
        [[x["equipment_id"],x["availability_pct"],x["failure_count"],x["unplanned_downtime_hours"],x["planned_downtime_hours"],x["mttr_hours"],x["mtbf_hours"],
          x["incidents_period"],x["open_incidents"],x["critical_open"],x["active_alarms"],x["current_state"]] for x in data["tool_matrix"]])
    ws=wb.create_sheet("Alarm Pareto");_sheet(ws,["Alarm Code","Message","Count"],[[x["alarm_code"],x["message"],x["count"]] for x in data["alarm_pareto"]])
    ws=wb.create_sheet("Incident Pareto");_sheet(ws,["Equipment","Incidents"],[[x["equipment_id"],x["count"]] for x in data["incident_pareto"]])
    ws=wb.create_sheet("Critical Incidents");_sheet(ws,["Ticket","Equipment","Priority","Status","Owner","Title","Description","Created","Updated"],
        [[x.ticket_no,x.equipment_id,x.priority,x.status,x.owner,x.title,x.description,x.created_at,x.updated_at] for x in metrics["critical_incidents"]])
    ws=wb.create_sheet("PM Exceptions");_sheet(ws,["Task","Equipment","PM","Name","Due","Planned","Status","Assigned","Hours","Priority"],
        [[x.id,x.equipment_id,x.pm_id,x.pm_name,x.original_due_date,x.scheduled_date,x.status,x.assigned_to,x.estimated_hours,x.priority] for x in metrics["overdue_pm"]])
    ws=wb.create_sheet("Priority Queue");_sheet(ws,["Severity","Type","Equipment","Key","Summary","Owner","Age h"],
        [[x.get("severity"),x.get("kind"),x.get("equipment_id"),x.get("key"),x.get("summary"),x.get("owner"),x.get("age_hours")] for x in metrics["attention"]])
    Path(path).parent.mkdir(parents=True,exist_ok=True);wb.save(path);return str(path)


def _work_order_closeout_context(db,work_order_no: str) -> dict[str,Any]:
    wo=db.get_work_order(work_order_no)
    if not wo:raise ValueError("Work order not found")
    close=db.work_order_closeout_status(work_order_no)
    links=db.list_work_order_links(work_order_no)
    events=db.list_work_order_events(work_order_no)
    logs=[x for x in db.list_work_logs(wo.equipment_id,False,2000) if x.entity_type=="WORK_ORDER" and x.entity_key==work_order_no]
    linked_ticket_keys=[x.entity_key for x in links if x.entity_type=="TICKET"]
    if wo.source_type=="TICKET" and wo.source_key and wo.source_key not in linked_ticket_keys:linked_ticket_keys.insert(0,wo.source_key)
    tickets=[x for x in db.list_tickets() if x.ticket_no in linked_ticket_keys]
    linked_qual_keys={x.entity_key for x in links if x.entity_type=="QUALIFICATION"}
    qualifications=[x for x in db.list_qualification_runs(wo.equipment_id) if x.run_no in linked_qual_keys]
    release_ids=set()
    for link in links:
        if link.entity_type=="RELEASE":
            try:release_ids.add(int(link.entity_key))
            except Exception:pass
    releases=[x for x in db.list_release_requests() if x.id in release_ids]
    reservations=[x for x in db.list_reservations() if close.get("source_pm_task_id") and x.pm_task_id==close["source_pm_task_id"]]
    attachments=list(db.list_attachments("WORK_ORDER",work_order_no))
    for ticket in tickets:attachments.extend(db.list_attachments("TICKET",ticket.ticket_no))
    for run in qualifications:attachments.extend(db.list_attachments("QUALIFICATION",run.run_no))
    for release in releases:attachments.extend(db.list_attachments("RELEASE",str(release.id)))
    seen=set();dedup=[]
    for row in attachments:
        key=(row.file_sha256,row.stored_path)
        if key in seen:continue
        seen.add(key);dedup.append(row)
    return {
        "work_order":wo,"closeout":close,"links":links,"events":events,"logs":logs,
        "tickets":tickets,"qualifications":qualifications,"releases":releases,
        "reservations":reservations,"attachments":dedup,
    }


def export_work_order_closeout_pptx(db,work_order_no: str,path: str) -> str:
    ctx=_work_order_closeout_context(db,work_order_no);wo=ctx["work_order"];close=ctx["closeout"]
    prs=Presentation()
    slide=prs.slides.add_slide(prs.slide_layouts[0]);slide.shapes.title.text=f"Return-to-Service Packet — {wo.work_order_no}"
    slide.placeholders[1].text=f"{wo.equipment_id} | {wo.title} | {wo.status} | Owner: {wo.owner or '—'}"

    slide=prs.slides.add_slide(prs.slide_layouts[1]);_add_bullets(slide,"Work / Closeout Summary",[
        f"Source: {wo.source_type}:{wo.source_key or '—'}",
        f"Scope: {wo.description or '—'}",
        f"Qualification required: {'Yes' if wo.qualification_required else 'No'} · Valid qualification: {close.get('valid_qualification_run') or 'None'}",
        f"Release required: {'Yes' if wo.release_required else 'No'} · Active release: {close.get('active_release_status') or 'None'}",
        f"Labor entries: {close.get('labor_entries',0)} · Evidence files: {close.get('attachment_count',0)}",
        f"Part reservations: {close.get('part_reservations',0)} · Active reservations: {close.get('active_part_reservations',0)}",
        f"Closeout blockers: {'; '.join(close.get('blockers',[])) or 'None'}",
    ])

    slide=prs.slides.add_slide(prs.slide_layouts[5]);_add_table(slide,"Linked Incident / Problem Context",
        ["Ticket","Priority","Status","Owner","Title","Root Cause","Corrective Action"],
        [[x.ticket_no,x.priority,x.status,x.owner,x.title,x.root_cause,x.corrective_action] for x in ctx["tickets"]],10)

    slide=prs.slides.add_slide(prs.slide_layouts[5]);_add_table(slide,"Work Order Lifecycle / Labor",
        ["Type","From / User","To / Start","Reason / End","Owner / Minutes","By / Status"],
        [
            *[["Lifecycle",x.from_state,x.to_state,x.reason,x.owner,x.changed_by] for x in ctx["events"]],
            *[["Labor",x.username,x.started_at,x.ended_at,x.duration_minutes,x.status] for x in ctx["logs"]],
        ],16)

    qual_rows=[]
    for run in ctx["qualifications"]:
        qual_rows.append([run.run_no,run.protocol_id,run.protocol_revision,run.status,run.started_by,run.verified_by,run.approved_by,run.expires_at])
    slide=prs.slides.add_slide(prs.slide_layouts[5]);_add_table(slide,"Qualification / Verification",
        ["Run","Protocol","Rev","Status","Started By","Verified By","Approved By","Expires"],qual_rows,12)

    release_rows=[]
    for rel in ctx["releases"]:
        checks=json.loads(rel.checks_json or "{}")
        release_rows.append([rel.id,rel.status,rel.related_ticket,rel.requested_by,rel.verified_by,rel.approved_by,
                             sum(1 for v in checks.values() if v),len(checks)])
    slide=prs.slides.add_slide(prs.slide_layouts[5]);_add_table(slide,"Release Control",
        ["Release","Status","Ticket","Requested By","Verified By","Approved By","Checks Pass","Checks Total"],release_rows,12)

    slide=prs.slides.add_slide(prs.slide_layouts[5]);_add_table(slide,"Parts / Reservations",
        ["Part","Location","Qty","Status","Reserved By","Reserved"],
        [[x.part_number,x.location_code,x.quantity,x.status,x.reserved_by,x.reserved_at] for x in ctx["reservations"]],14)

    slide=prs.slides.add_slide(prs.slide_layouts[5]);_add_table(slide,"Linked Record Traceability",
        ["Type","Key","Relation","Created By","Created"],
        [[x.entity_type,x.entity_key,x.relation,x.created_by,x.created_at] for x in ctx["links"]],16)

    _add_evidence_slides(prs,ctx["attachments"],"Closeout")
    Path(path).parent.mkdir(parents=True,exist_ok=True);prs.save(path);return str(path)


def export_work_order_closeout_xlsx(db,work_order_no: str,path: str) -> str:
    ctx=_work_order_closeout_context(db,work_order_no);wo=ctx["work_order"];close=ctx["closeout"]
    wb=Workbook();summary=wb.active;summary.title="Closeout Summary"
    _sheet(summary,["Field","Value"],[
        ["Work Order",wo.work_order_no],["Equipment",wo.equipment_id],["Title",wo.title],["Status",wo.status],
        ["Priority",wo.priority],["Owner",wo.owner],["Team",wo.team],["Source",f"{wo.source_type}:{wo.source_key}"],
        ["Scope",wo.description],["Qualification Required",wo.qualification_required],
        ["Valid Qualification",close.get("valid_qualification_run","")],["Open Qualification",close.get("open_qualification_run","")],
        ["Release Required",wo.release_required],["Active Release ID",close.get("active_release_id","")],
        ["Active Release Status",close.get("active_release_status","")],
        ["Closeout Blockers","; ".join(close.get("blockers",[]))],
    ])
    ws=wb.create_sheet("Incidents");_sheet(ws,["Ticket","Priority","Status","Owner","Title","Problem","Root Cause","Corrective Action","Verification"],
        [[x.ticket_no,x.priority,x.status,x.owner,x.title,x.description,x.root_cause,x.corrective_action,x.verification] for x in ctx["tickets"]])
    ws=wb.create_sheet("Work Lifecycle");_sheet(ws,["From","To","Reason","Owner","Changed By","Time"],
        [[x.from_state,x.to_state,x.reason,x.owner,x.changed_by,x.occurred_at] for x in ctx["events"]])
    ws=wb.create_sheet("Labor");_sheet(ws,["User","Type","Start","End","Minutes","Status","Note"],
        [[x.username,x.work_type,x.started_at,x.ended_at,x.duration_minutes,x.status,x.note] for x in ctx["logs"]])
    qrows=[]
    for run in ctx["qualifications"]:
        qrows.append([run.run_no,run.protocol_id,run.protocol_revision,run.protocol_name,run.status,run.started_by,run.started_at,run.submitted_by,run.submitted_at,run.verified_by,run.verified_at,run.approved_by,run.approved_at,run.expires_at,run.conclusion])
    ws=wb.create_sheet("Qualifications");_sheet(ws,["Run","Protocol","Rev","Name","Status","Started By","Started","Submitted By","Submitted","Verified By","Verified","Approved By","Approved","Expires","Conclusion"],qrows)
    check_rows=[]
    for run in ctx["qualifications"]:
        checks=json.loads(run.frozen_checks_json or "[]");results=json.loads(run.results_json or "{}")
        for check in checks:
            cid=str(check.get("check_id") or check.get("id") or "")
            result=results.get(cid,{}) if isinstance(results,dict) else {}
            check_rows.append([run.run_no,cid,check.get("label") or check.get("name") or "",check.get("acceptance",""),result.get("result",""),result.get("comment",""),result.get("entered_by",""),result.get("entered_at","")])
    ws=wb.create_sheet("Qualification Checks");_sheet(ws,["Run","Check","Description","Acceptance","Result","Comment","Entered By","Entered"],check_rows)
    relrows=[];checklist=[]
    for rel in ctx["releases"]:
        relrows.append([rel.id,rel.status,rel.related_ticket,rel.requested_by,rel.requested_at,rel.verified_by,rel.verified_at,rel.approved_by,rel.approved_at,rel.notes])
        for key,value in json.loads(rel.checks_json or "{}").items():checklist.append([rel.id,key,value])
    ws=wb.create_sheet("Releases");_sheet(ws,["ID","Status","Ticket","Requested By","Requested","Verified By","Verified","Approved By","Approved","Notes"],relrows)
    ws=wb.create_sheet("Release Checklist");_sheet(ws,["Release","Check","Pass"],checklist)
    ws=wb.create_sheet("Parts");_sheet(ws,["Part","Location","Qty","Status","Reserved By","Reserved","Released"],
        [[x.part_number,x.location_code,x.quantity,x.status,x.reserved_by,x.reserved_at,x.released_at] for x in ctx["reservations"]])
    ws=wb.create_sheet("Links");_sheet(ws,["Type","Key","Relation","Created By","Created"],
        [[x.entity_type,x.entity_key,x.relation,x.created_by,x.created_at] for x in ctx["links"]])
    ws=wb.create_sheet("Evidence");_sheet(ws,["Entity","Name","Category","Caption","Tags","Path","SHA256","Added By","Added"],
        [[f"{x.entity_type}:{x.entity_key}",x.original_name,x.category,x.caption,x.tags,x.stored_path,x.file_sha256,x.created_by,x.created_at] for x in ctx["attachments"]])
    Path(path).parent.mkdir(parents=True,exist_ok=True);wb.save(path);return str(path)
