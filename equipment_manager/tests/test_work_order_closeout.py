import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation

from database import Database
from reporting import export_work_order_closeout_pptx, export_work_order_closeout_xlsx


class WorkOrderCloseoutTests(unittest.TestCase):
    def setUp(self):
        self.db=Database("sqlite:///:memory:")
        self.db.save_equipment({
            "equipment_id":"ETCH-CLOSE","name":"Closeout Etcher","equipment_type":"Plasma Etcher",
        },user="seed")
        self.db.save_qualification_protocol(
            "QUAL-CLOSE","Controlled return-to-service",
            [{"check_id":"Q01","label":"Verification wafer","acceptance":"Pass"}],
            "author",equipment_type="Plasma Etcher",
        )
        self.wo=self.db.create_work_order({
            "equipment_id":"ETCH-CLOSE","source_type":"ENGINEERING","title":"Major chamber repair",
            "owner":"executor","priority":"High","qualification_required":True,"release_required":True,
        },"creator")
        self.wo=self.db.transition_work_order(self.wo.work_order_no,"In Progress","executor",expected_version=self.wo.version)
        self.wo=self.db.transition_work_order(self.wo.work_order_no,"Ready for Qualification","executor",expected_version=self.wo.version)

    def test_closeout_packet_drives_qualification_then_release_request(self):
        before=self.db.work_order_closeout_status(self.wo.work_order_no)
        self.assertTrue(before["can_start_qualification"])
        self.assertFalse(before["can_request_release"])
        self.assertIn("Approved valid qualification is required.",before["blockers"])

        run=self.db.start_work_order_qualification(self.wo.work_order_no,"executor")
        links=self.db.list_work_order_links(self.wo.work_order_no)
        self.assertTrue(any(x.entity_type=="QUALIFICATION" and x.entity_key==run.run_no for x in links))

        run=self.db.save_qualification_result(run.id,"Q01","PASS","Passed","executor",expected_version=run.version)
        run=self.db.submit_qualification_run(run.id,"executor","All checks pass",expected_version=run.version)
        run=self.db.verify_qualification_run(run.id,"verifier","Independent review",expected_version=run.version)
        run=self.db.approve_qualification_run(run.id,"manager",30,"Approved",expected_version=run.version)

        ready=self.db.work_order_closeout_status(self.wo.work_order_no)
        self.assertEqual(ready["valid_qualification_run"],run.run_no)
        self.assertTrue(ready["can_request_release"])

        release=self.db.create_work_order_release_request(self.wo.work_order_no,"release_requester")
        checks=json.loads(release.checks_json)
        self.assertTrue(checks)
        self.assertFalse(any(checks.values()))
        links=self.db.list_work_order_links(self.wo.work_order_no)
        self.assertTrue(any(x.entity_type=="RELEASE" and x.entity_key==str(release.id) for x in links))

        after=self.db.work_order_closeout_status(self.wo.work_order_no)
        self.assertEqual(after["active_release_id"],release.id)
        self.assertFalse(after["can_request_release"])

    def test_closeout_review_pack_contains_qualification_release_and_traceability(self):
        run=self.db.start_work_order_qualification(self.wo.work_order_no,"executor")
        run=self.db.save_qualification_result(run.id,"Q01","PASS","Passed","executor",expected_version=run.version)
        run=self.db.submit_qualification_run(run.id,"executor","All checks pass",expected_version=run.version)
        run=self.db.verify_qualification_run(run.id,"verifier","Independent review",expected_version=run.version)
        run=self.db.approve_qualification_run(run.id,"manager",30,"Approved",expected_version=run.version)
        release=self.db.create_work_order_release_request(self.wo.work_order_no,"release_requester")
        with tempfile.TemporaryDirectory() as td:
            ppt=Path(td)/"closeout.pptx";xlsx=Path(td)/"closeout.xlsx"
            export_work_order_closeout_pptx(self.db,self.wo.work_order_no,str(ppt))
            export_work_order_closeout_xlsx(self.db,self.wo.work_order_no,str(xlsx))
            deck=Presentation(str(ppt))
            self.assertGreaterEqual(len(deck.slides),7)
            self.assertIn("Return-to-Service Packet",deck.slides[0].shapes.title.text)
            wb=load_workbook(str(xlsx),read_only=True)
            for sheet in ["Closeout Summary","Qualifications","Qualification Checks","Releases","Release Checklist","Links","Evidence"]:
                self.assertIn(sheet,wb.sheetnames)
            release_sheet=wb["Releases"]
            values=list(release_sheet.values)
            self.assertTrue(any(str(release.id)==str(row[0]) for row in values[1:]))

    def test_closeout_isolates_qualification_by_work_order(self):
        run=self.db.start_work_order_qualification(self.wo.work_order_no,"executor")
        run=self.db.save_qualification_result(run.id,"Q01","PASS","Passed","executor",expected_version=run.version)
        run=self.db.submit_qualification_run(run.id,"executor","All checks pass",expected_version=run.version)
        run=self.db.verify_qualification_run(run.id,"verifier","Independent review",expected_version=run.version)
        run=self.db.approve_qualification_run(run.id,"manager",30,"Approved",expected_version=run.version)

        other=self.db.create_work_order({
            "equipment_id":"ETCH-CLOSE","source_type":"ENGINEERING","title":"Second major repair",
            "owner":"executor","qualification_required":True,"release_required":True,
        },"creator")
        other=self.db.transition_work_order(other.work_order_no,"In Progress","executor",expected_version=other.version)
        other=self.db.transition_work_order(other.work_order_no,"Ready for Qualification","executor",expected_version=other.version)

        state=self.db.work_order_closeout_status(other.work_order_no)
        self.assertEqual(state["valid_qualification_run"],"")
        self.assertTrue(state["can_start_qualification"])
        self.assertIn("Approved valid qualification is required.",state["blockers"])

    def test_closeout_does_not_bypass_work_completion_state(self):
        other=self.db.create_work_order({
            "equipment_id":"ETCH-CLOSE","source_type":"ENGINEERING","title":"Incomplete work",
            "owner":"executor","qualification_required":True,"release_required":True,
        },"creator")
        state=self.db.work_order_closeout_status(other.work_order_no)
        self.assertFalse(state["can_start_qualification"])
        self.assertTrue(any("still Open" in x for x in state["blockers"]))
        with self.assertRaises(ValueError):
            self.db.start_work_order_qualification(other.work_order_no,"executor")


if __name__=="__main__":
    unittest.main()
