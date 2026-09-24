import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation

from database import Database
from reporting import (
    export_equipment_pptx,export_equipment_xlsx,
    export_incident_pptx,export_incident_xlsx,
    export_pm_execution_pptx,export_pm_execution_xlsx,
    export_work_order_pptx,export_work_order_xlsx,
    export_qualification_pptx,export_qualification_xlsx,
    export_release_pptx,export_release_xlsx,
    export_engineering_review_pptx,export_engineering_review_xlsx,
)


class OfficeReportingTests(unittest.TestCase):
    def setUp(self):
        self.db=Database("sqlite:///:memory:")
        self.db.create_user("ee","Equipment Engineer","engineer-password-123","Equipment Engineer")
        self.db.save_equipment({
            "equipment_id":"ETCH-RPT","name":"Etcher Reporting","equipment_type":"Etch",
            "area":"ETCH","model":"RPT-100","serial_number":"SN-RPT","owner":"ee",
        },user="ee")
        self.ticket=self.db.save_ticket({
            "ticket_no":"INC-RPT-1","equipment_id":"ETCH-RPT","title":"Pressure instability",
            "description":"Pressure oscillation during process","severity":"S2","priority":"P2","owner":"ee",
            "root_cause":"Pending","corrective_action":"Inspect pressure loop","verification":"Pending","created_by":"ee",
        })
        self.db.save_ticket_operational_control(
            self.ticket.ticket_no,
            {"containment":"Tool held","production_impact":"Tool unavailable","affected_lots":"LOT-1",
             "safety_quality_risk":"Potential process drift","response_due_at":None,
             "containment_due_at":None,"resolution_due_at":None},
            "ee",
        )
        self.db.save_incident_why(self.ticket.ticket_no,1,"Why did pressure oscillate?","Control response unstable","ee")
        self.db.save_incident_causal_factor(
            self.ticket.ticket_no,
            {"category":"Machine","factor_type":"Suspected","description":"Pressure controller","evidence":"Trend correlation","status":"Open"},
            "ee",
        )
        self.db.save_incident_action(
            self.ticket.ticket_no,
            {"action_type":"Corrective","description":"Inspect controller","owner":"ee","due_at":None,"effectiveness_criteria":"Stable pressure"},
            "ee",
        )
        self.db.upsert_pm_spec({
            "pm_id":"PM-RPT","step_no":1,"activity":"Verify chamber pressure","method":"Gauge",
            "input_type":"Numeric","unit":"Pa","spec_low":0.0,"spec_high":10.0,
            "acceptance_text":"0–10 Pa","reaction_plan":"Stop and investigate","revision":1,"active":True,
        })
        self.pm_task=self.db.upsert_pm_task({
            "equipment_id":"ETCH-RPT","pm_id":"PM-RPT","pm_name":"Reporting PM",
            "status":"Scheduled","assigned_to":"ee","estimated_hours":1.5,"priority":"Normal",
        })
        self.pm_execution=self.db.start_pm_execution(self.pm_task.id,"ee")
        self.db.save_pm_result(self.pm_execution.id,1,{
            "value_text":"5.0","value_numeric":5.0,"result":"PASS","comment":"Stable","entered_by":"ee",
        })
        self.work_order=self.db.create_work_order_from_ticket(self.ticket.ticket_no,"ee")
        self.protocol=self.db.save_qualification_protocol(
            "QUAL-RPT","Reporting qualification",
            [{"check_id":"Q01","label":"Monitor result","acceptance":"Pass"}],
            "ee",equipment_id="ETCH-RPT",
        )
        self.qual_run=self.db.start_qualification_run("ETCH-RPT","QUAL-RPT","ee")
        self.qual_run=self.db.save_qualification_result(
            self.qual_run.id,"Q01","PASS","Monitor passed","ee",expected_version=self.qual_run.version,
        )
        checks={
            "maintenance_complete":True,"measurements_pass":True,"calibration_valid":True,
            "safety_check":True,"verification_run":True,"critical_tickets_cleared":True,
        }
        self.release=self.db.create_release_request("ETCH-RPT",self.ticket.ticket_no,checks,"Review packet","ee")

    def test_incident_pptx_and_xlsx_are_reopenable(self):
        with tempfile.TemporaryDirectory() as td:
            ppt=Path(td)/"incident.pptx";xlsx=Path(td)/"incident.xlsx"
            export_incident_pptx(self.db,self.ticket.ticket_no,str(ppt))
            export_incident_xlsx(self.db,self.ticket.ticket_no,str(xlsx))
            deck=Presentation(str(ppt))
            self.assertGreaterEqual(len(deck.slides),6)
            wb=load_workbook(str(xlsx),read_only=True)
            self.assertIn("Summary",wb.sheetnames)
            self.assertIn("5-Why",wb.sheetnames)
            self.assertIn("CAPA",wb.sheetnames)

    def test_pm_execution_pptx_and_xlsx_are_reopenable(self):
        with tempfile.TemporaryDirectory() as td:
            ppt=Path(td)/"pm.pptx";xlsx=Path(td)/"pm.xlsx"
            export_pm_execution_pptx(self.db,self.pm_task.id,str(ppt))
            export_pm_execution_xlsx(self.db,self.pm_task.id,str(xlsx))
            deck=Presentation(str(ppt));self.assertGreaterEqual(len(deck.slides),3)
            wb=load_workbook(str(xlsx),read_only=True)
            self.assertIn("Summary",wb.sheetnames);self.assertIn("Checklist Results",wb.sheetnames)
            self.assertIn("Requirements",wb.sheetnames)

    def test_work_order_pptx_and_xlsx_are_reopenable(self):
        with tempfile.TemporaryDirectory() as td:
            ppt=Path(td)/"wo.pptx";xlsx=Path(td)/"wo.xlsx"
            export_work_order_pptx(self.db,self.work_order.work_order_no,str(ppt))
            export_work_order_xlsx(self.db,self.work_order.work_order_no,str(xlsx))
            deck=Presentation(str(ppt));self.assertGreaterEqual(len(deck.slides),5)
            wb=load_workbook(str(xlsx),read_only=True)
            self.assertIn("Summary",wb.sheetnames);self.assertIn("Lifecycle",wb.sheetnames)
            self.assertIn("Links",wb.sheetnames)

    def test_qualification_and_release_report_packs_are_reopenable(self):
        with tempfile.TemporaryDirectory() as td:
            qp=Path(td)/"qual.pptx";qx=Path(td)/"qual.xlsx"
            export_qualification_pptx(self.db,self.qual_run.run_no,str(qp))
            export_qualification_xlsx(self.db,self.qual_run.run_no,str(qx))
            self.assertGreaterEqual(len(Presentation(str(qp)).slides),3)
            qwb=load_workbook(str(qx),read_only=True);self.assertIn("Checks",qwb.sheetnames)

            rp=Path(td)/"release.pptx";rx=Path(td)/"release.xlsx"
            export_release_pptx(self.db,self.release.id,str(rp))
            export_release_xlsx(self.db,self.release.id,str(rx))
            self.assertGreaterEqual(len(Presentation(str(rp)).slides),3)
            rwb=load_workbook(str(rx),read_only=True);self.assertIn("Checklist",rwb.sheetnames)

    def test_engineering_review_pptx_and_xlsx_are_reopenable(self):
        with tempfile.TemporaryDirectory() as td:
            ppt=Path(td)/"review.pptx";xlsx=Path(td)/"review.xlsx"
            export_engineering_review_pptx(self.db,30,str(ppt))
            export_engineering_review_xlsx(self.db,30,str(xlsx))
            deck=Presentation(str(ppt));self.assertGreaterEqual(len(deck.slides),6)
            wb=load_workbook(str(xlsx),read_only=True)
            self.assertIn("Summary",wb.sheetnames);self.assertIn("Tool Matrix",wb.sheetnames)
            self.assertIn("Alarm Pareto",wb.sheetnames);self.assertIn("Trend",wb.sheetnames)

    def test_equipment_pptx_and_xlsx_are_reopenable(self):
        with tempfile.TemporaryDirectory() as td:
            ppt=Path(td)/"equipment.pptx";xlsx=Path(td)/"equipment.xlsx"
            export_equipment_pptx(self.db,"ETCH-RPT",str(ppt))
            export_equipment_xlsx(self.db,"ETCH-RPT",str(xlsx))
            deck=Presentation(str(ppt))
            self.assertGreaterEqual(len(deck.slides),5)
            wb=load_workbook(str(xlsx),read_only=True)
            self.assertIn("Summary",wb.sheetnames)
            self.assertIn("Activity",wb.sheetnames)
            self.assertIn("Incidents",wb.sheetnames)


if __name__=="__main__":
    unittest.main()
