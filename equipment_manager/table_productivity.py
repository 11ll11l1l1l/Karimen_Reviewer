from __future__ import annotations

from datetime import datetime
import csv
import hashlib
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from PySide6.QtCore import Qt
from PySide6.QtGui import QAction, QKeySequence, QShortcut
from PySide6.QtWidgets import QApplication, QFileDialog, QMenu, QMessageBox, QTableWidget



_PREF_DB=None
_PREF_USER=""


def configure_productivity_context(db,username: str):
    global _PREF_DB,_PREF_USER
    _PREF_DB=db;_PREF_USER=username or ""


def _view_key(table: QTableWidget) -> str:
    raw="|".join(_headers(table))
    return "table_view."+hashlib.sha1(raw.encode("utf-8")).hexdigest()[:16]


def _save_view(table: QTableWidget):
    if not _PREF_DB or not _PREF_USER:return
    header=table.horizontalHeader()
    value={
        "hidden":[col for col in range(table.columnCount()) if table.isColumnHidden(col)],
        "visual_order":[header.logicalIndex(v) for v in range(header.count())],
    }
    _PREF_DB.set_user_preference(_PREF_USER,_view_key(table),value)


def _restore_view(table: QTableWidget):
    if not _PREF_DB or not _PREF_USER:return
    value=_PREF_DB.get_user_preference(_PREF_USER,_view_key(table),{})
    if not isinstance(value,dict):return
    hidden={int(x) for x in value.get("hidden",[]) if isinstance(x,int) or str(x).isdigit()}
    for col in range(table.columnCount()):table.setColumnHidden(col,col in hidden)
    order=value.get("visual_order",[])
    header=table.horizontalHeader()
    if isinstance(order,list) and len(order)==header.count():
        try:
            for target,logical in enumerate(order):
                logical=int(logical);current=header.visualIndex(logical)
                if current>=0 and current!=target:header.moveSection(current,target)
        except Exception:pass


def _headers(table: QTableWidget) -> list[str]:
    out=[]
    for c in range(table.columnCount()):
        item=table.horizontalHeaderItem(c)
        out.append(item.text() if item else f"Column {c+1}")
    return out


def _cell(table: QTableWidget,row: int,col: int) -> str:
    item=table.item(row,col)
    return item.text() if item else ""


def copy_selection_tsv(table: QTableWidget):
    ranges=table.selectedRanges()
    if not ranges:return
    r0=min(x.topRow() for x in ranges);r1=max(x.bottomRow() for x in ranges)
    c0=min(x.leftColumn() for x in ranges);c1=max(x.rightColumn() for x in ranges)
    selected={(idx.row(),idx.column()) for idx in table.selectedIndexes()}
    lines=[]
    for r in range(r0,r1+1):
        values=[]
        for c in range(c0,c1+1):
            values.append(_cell(table,r,c) if (r,c) in selected else "")
        lines.append("\t".join(values))
    QApplication.clipboard().setText("\n".join(lines))


def copy_rows_with_headers(table: QTableWidget,selected_only: bool=False):
    rows=sorted({idx.row() for idx in table.selectedIndexes()}) if selected_only else list(range(table.rowCount()))
    if not rows:return
    columns=[c for c in range(table.columnCount()) if not table.isColumnHidden(c)]
    headers=_headers(table)
    lines=["\t".join(headers[c] for c in columns)]
    for r in rows:lines.append("\t".join(_cell(table,r,c) for c in columns))
    QApplication.clipboard().setText("\n".join(lines))


def export_table_csv(table: QTableWidget,selected_only: bool=False,parent=None):
    rows=sorted({idx.row() for idx in table.selectedIndexes()}) if selected_only else list(range(table.rowCount()))
    if not rows:
        QMessageBox.information(parent or table,"Export","No rows selected.")
        return None
    suggested=table.property("ems_export_name") or "EMS_Export"
    safe="".join(ch if ch.isalnum() or ch in {"-","_"} else "_" for ch in str(suggested)).strip("_") or "EMS_Export"
    default=str(Path.cwd()/f"{safe}_{datetime.now():%Y%m%d_%H%M}.csv")
    path,_=QFileDialog.getSaveFileName(parent or table,"Export to CSV",default,"CSV (*.csv)")
    if not path:return None
    if not path.lower().endswith(".csv"):path+=".csv"
    columns=[c for c in range(table.columnCount()) if not table.isColumnHidden(c)]
    try:
        with open(path,"w",encoding="utf-8-sig",newline="") as handle:
            writer=csv.writer(handle)
            headers=_headers(table);writer.writerow([headers[c] for c in columns])
            for r in rows:writer.writerow([_cell(table,r,c) for c in columns])
    except Exception as exc:
        QMessageBox.critical(parent or table,"CSV export",str(exc));return None
    QMessageBox.information(parent or table,"CSV export",f"Exported {len(rows)} row(s).\n{path}")
    return path


def _write_workbook(table: QTableWidget,path: str,rows: list[int]):
    wb=Workbook();ws=wb.active;ws.title="Export"
    headers=_headers(table)
    for c,label in enumerate(headers,1):
        cell=ws.cell(1,c,label);cell.font=Font(bold=True);cell.fill=PatternFill("solid",fgColor="D9E2F3")
    for out_r,source_r in enumerate(rows,2):
        for c in range(table.columnCount()):
            ws.cell(out_r,c+1,_cell(table,source_r,c))
    ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions
    for c,label in enumerate(headers,1):
        width=max(len(label),*(len(str(ws.cell(r,c).value or "")) for r in range(2,min(ws.max_row,250)+1)))
        ws.column_dimensions[get_column_letter(c)].width=min(max(width+2,10),60)
    meta=wb.create_sheet("_EMS Export")
    meta["A1"]="Generated";meta["B1"]=datetime.now().isoformat(timespec="seconds")
    meta["A2"]="Rows";meta["B2"]=len(rows)
    meta.sheet_state="hidden"
    wb.save(path)


def export_table_xlsx(table: QTableWidget,selected_only: bool=False,parent=None):
    rows=sorted({idx.row() for idx in table.selectedIndexes()}) if selected_only else list(range(table.rowCount()))
    if not rows:
        QMessageBox.information(parent or table,"Export","No rows selected.")
        return None
    suggested=table.property("ems_export_name") or "EMS_Export"
    safe="".join(ch if ch.isalnum() or ch in {"-","_"} else "_" for ch in str(suggested)).strip("_") or "EMS_Export"
    default=str(Path.cwd()/f"{safe}_{datetime.now():%Y%m%d_%H%M}.xlsx")
    path,_=QFileDialog.getSaveFileName(parent or table,"Export to Excel",default,"Excel Workbook (*.xlsx)")
    if not path:return None
    if not path.lower().endswith(".xlsx"):path+=".xlsx"
    try:_write_workbook(table,path,rows)
    except Exception as exc:
        QMessageBox.critical(parent or table,"Excel export",str(exc));return None
    QMessageBox.information(parent or table,"Excel export",f"Exported {len(rows)} row(s).\n{path}")
    return path


def install_table_productivity(table: QTableWidget,export_name: str="EMS Export"):
    if table.property("ems_productivity_installed"):return table
    table.setProperty("ems_productivity_installed",True)
    table.setProperty("ems_export_name",export_name)
    table.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)

    header=table.horizontalHeader();header.setSectionsMovable(True)
    def menu_at(pos):
        menu=QMenu(table)
        copy_action=menu.addAction("Copy selected cells")
        copy_rows=menu.addAction("Copy selected rows with headers")
        copy_all=menu.addAction("Copy current table with headers")
        menu.addSeparator()
        export_selected=menu.addAction("Export selected rows to Excel")
        export_all=menu.addAction("Export current table to Excel")
        csv_selected=menu.addAction("Export selected rows to CSV")
        csv_all=menu.addAction("Export current table to CSV")
        menu.addSeparator()
        columns=menu.addMenu("Columns")
        column_actions={}
        for col,label in enumerate(_headers(table)):
            action=columns.addAction(label);action.setCheckable(True);action.setChecked(not table.isColumnHidden(col));column_actions[action]=col
        save_view=menu.addAction("Save table view")
        reset_view=menu.addAction("Reset table view")
        action=menu.exec(table.viewport().mapToGlobal(pos))
        if action==copy_action:copy_selection_tsv(table)
        elif action==copy_rows:copy_rows_with_headers(table,True)
        elif action==copy_all:copy_rows_with_headers(table,False)
        elif action==export_selected:export_table_xlsx(table,True,table)
        elif action==export_all:export_table_xlsx(table,False,table)
        elif action==csv_selected:export_table_csv(table,True,table)
        elif action==csv_all:export_table_csv(table,False,table)
        elif action==save_view:_save_view(table)
        elif action==reset_view:
            for col in range(table.columnCount()):table.setColumnHidden(col,False)
            if _PREF_DB and _PREF_USER:_PREF_DB.set_user_preference(_PREF_USER,_view_key(table),{})
        elif action in column_actions:
            col=column_actions[action];table.setColumnHidden(col,not action.isChecked());_save_view(table)

    table.customContextMenuRequested.connect(menu_at)
    _restore_view(table)
    shortcut=QShortcut(QKeySequence.Copy,table)
    shortcut.activated.connect(lambda:copy_selection_tsv(table))
    table._ems_copy_shortcut=shortcut
    return table
